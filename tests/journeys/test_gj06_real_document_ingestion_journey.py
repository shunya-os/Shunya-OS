"""GJ-06 — REAL document ingestion (PDF / DOCX / XLSX), driven over real HTTP.

§8 of the directive requires document intelligence to work on REPRESENTATIVE
REAL documents whose organisation is derived from CONTENT UNDERSTANDING, not from
the filename or the folder path.

To make that claim falsifiable this journey deliberately uploads files with
NEUTRAL filenames (`doc-8371.pdf`, `scan-2210.docx`, ...). If the itinerary
hierarchy came from the name it could not appear; the only source of "Bali" and
"International" is the text extracted from inside the real PDF bytes.

Documents are GENERATED at test time (reportlab / python-docx / openpyxl), so no
binary fixtures are committed and the bytes are genuinely parsable by the same
libraries production uses (pdfplumber / python-docx / openpyxl).

IMAGE + OCR: an image cannot be content-classified without an OCR provider.
pytesseract is NOT installed in this environment, so the image case asserts only
that ingestion is honest about it — it is recorded, not claimed as working.
"""
from __future__ import annotations

import http.cookiejar
import io
import json
import os
import tempfile
import threading
import urllib.error
import urllib.request
import uuid

import pytest
from werkzeug.serving import make_server

ORG_ID = 981
ALT_ORG_ID = 982
EMAIL = "gj06-real-docs@example.com"
EMAIL_ALT = "gj06-real-docs-alt@example.com"
PASSWORD = "gj06-real-docs-pass"

SHELL_MARKER = "SHUNYA_GJ06_REAL_DOCS_SHELL"
SHELL_HTML = (
    "<!doctype html><html><head><title>SHUNYA</title></head>"
    f"<body>{SHELL_MARKER}</body></html>"
)

# Deliberately neutral: nothing here hints at itinerary/Bali/International.
NEUTRAL_PDF_NAME = "doc-8371.pdf"
NEUTRAL_DOCX_NAME = "scan-2210.docx"
NEUTRAL_XLSX_NAME = "sheet-9930.xlsx"
NEUTRAL_PNG_NAME = "image-4471.png"

BALI_ITINERARY_TEXT = [
    "Travel Itinerary",
    "Passenger: Nishesh Sharma    Booking reference: SHY-4N5D-2026",
    "",
    "Trip: Bali 4N5D",
    "Destination: Bali, Indonesia",
    "Departure: 2026-11-02    Arrival: 2026-11-02",
    "Check-in: 2026-11-02    Check-out: 2026-11-06",
    "",
    "Day 1  Arrival in Denpasar, hotel transfer, welcome dinner",
    "Day 2  Ubud rice terraces and temple visit",
    "Day 3  Nusa Penida day trip",
    "Day 4  Spa retreat and beach day",
    "Day 5  Departure",
    "",
    "Passport and visa details held on file.",
]

QUOTATION_TEXT = [
    "QUOTATION",
    "Quotation number: QT-2026-0417",
    "Vendor: Bali Luxury Transfers",
    "Valid until: 2026-12-31",
    "",
    "Unit price: 4500 per vehicle transfer",
    "Estimate for 4 transfers: 18000",
    "Subtotal: 18000",
    "Payment terms: 30 days from invoice",
    "",
    "This quotation is an offer subject to availability.",
]


# ---------------------------------------------------------------------------
# Real document generators
# ---------------------------------------------------------------------------

def make_pdf(lines: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    _, height = A4
    y = height - 60
    for line in lines:
        c.drawString(55, y, line)
        y -= 16
    c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(lines: list[str]) -> bytes:
    import docx

    d = docx.Document()
    for line in lines:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Destination", "Nights", "Guests", "Total"])
    ws.append(["Bali", 4, 2, 18000])
    ws.append(["Phuket", 3, 2, 12000])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_png() -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (600, 200), "white")
    ImageDraw.Draw(img).text((10, 90), "Bali 4N5D itinerary", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# ---------------------------------------------------------------------------
# HTTP harness (multipart upload)
# ---------------------------------------------------------------------------

class Http:
    def __init__(self, base: str):
        self.base = base
        self.jar = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(
            urllib.request.HTTPCookieProcessor(self.jar))

    def _send(self, req):
        try:
            with self.opener.open(req, timeout=40) as resp:
                return resp.status, resp.read().decode()
        except urllib.error.HTTPError as exc:
            return exc.code, exc.read().decode()

    def json(self, method, path, body=None):
        data = json.dumps(body).encode() if body is not None else None
        headers = {"Content-Type": "application/json"} if body is not None else {}
        status, text = self._send(urllib.request.Request(
            f"{self.base}{path}", data=data, method=method, headers=headers))
        try:
            return status, json.loads(text)
        except json.JSONDecodeError:
            return status, {"_raw": text[:400]}

    def upload(self, path: str, filename: str, content: bytes, content_type: str):
        boundary = "----gj06" + uuid.uuid4().hex
        body = b"".join([
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'.encode(),
            f"Content-Type: {content_type}\r\n\r\n".encode(),
            content,
            f"\r\n--{boundary}--\r\n".encode(),
        ])
        req = urllib.request.Request(
            f"{self.base}{path}", data=body, method="POST",
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
        status, text = self._send(req)
        try:
            return status, json.loads(text)
        except json.JSONDecodeError:
            return status, {"_raw": text[:400]}


class Server:
    def __init__(self, app):
        self._server = make_server("127.0.0.1", 0, app, threaded=True)
        self.port = self._server.server_port
        self._thread = threading.Thread(target=self._server.serve_forever, daemon=True)

    @property
    def base(self):
        return f"http://127.0.0.1:{self.port}"

    def start(self):
        self._thread.start()
        return self

    def stop(self):
        self._server.shutdown()


@pytest.fixture(autouse=True)
def _point_frontend_at_release(tmp_path_factory, monkeypatch):
    dist = tmp_path_factory.mktemp("gj06_release")
    (dist / "index.html").write_text(SHELL_HTML, encoding="utf-8")
    monkeypatch.setenv("SHUNYA_FRONTEND_DIST", str(dist))


@pytest.fixture(scope="module")
def journey_app(tmp_path_factory):
    db_file = tmp_path_factory.mktemp("gj06_db") / "journey.db"
    from app import create_app, db

    app = create_app({
        "TESTING": True,
        "SQLALCHEMY_DATABASE_URI": f"sqlite:///{db_file}",
        "WTF_CSRF_ENABLED": False,
    })
    declared_secure = app.config["SESSION_COOKIE_SECURE"]
    app.config["SESSION_COOKIE_SECURE"] = False

    with app.app_context():
        db.create_all()
        from app.auth import TeamMember
        from app.authz.services import seed_default_roles
        from app.models import OrgMember, Organization
        from app.objects.legacy_models import ShWorkspaceMembership, Workspace

        for oid, name, slug, email, ws_id in (
            (ORG_ID, "GJ06 Org", "gj06", EMAIL, "ws_gj06_primary"),
            (ALT_ORG_ID, "GJ06 Other Org", "gj06-other", EMAIL_ALT, "ws_gj06_alt"),
        ):
            if not db.session.get(Organization, oid):
                db.session.add(Organization(id=oid, name=name, slug=slug,
                                            is_active=True))
                db.session.flush()
                seed_default_roles(oid)
            # A workspace is required: sign-in resolves the org context through
            # workspace membership, and without it every request would run with
            # no organization context.
            if not db.session.get(Workspace, ws_id):
                db.session.add(Workspace(id=ws_id, name=f"GJ06 WS {oid}",
                                         workspace_type="business", status="active",
                                         created_by="gj06_fixture",
                                         organization_id=oid))
                db.session.flush()
            tm = TeamMember.query.filter_by(email=email).first()
            if not tm:
                tm = TeamMember(name=f"GJ06 {oid}", email=email,
                                is_active=True, verified=True)
                tm.set_password(PASSWORD)
                db.session.add(tm)
                db.session.flush()
            identity_id = str(tm.id)
            if not OrgMember.query.filter_by(organization_id=oid,
                                             email=email).first():
                db.session.add(OrgMember(organization_id=oid,
                                         identity_id=identity_id,
                                         name=f"GJ06 {oid}", email=email,
                                         role="owner"))
            if not ShWorkspaceMembership.query.filter_by(
                    workspace_id=ws_id, identity_id=identity_id).first():
                db.session.add(ShWorkspaceMembership(workspace_id=ws_id,
                                                     identity_id=identity_id,
                                                     role="owner", is_active=True))
        db.session.commit()

    app.config["_JOURNEY_DECLARED_SECURE_COOKIE"] = declared_secure
    return app


@pytest.fixture(scope="module")
def server(journey_app):
    srv = Server(journey_app).start()
    yield srv
    srv.stop()


def _signin(http, email=EMAIL):
    status, body = http.json("POST", "/api/v1/founder/signin",
                             {"email": email, "password": PASSWORD})
    return body.get("identity_id", "") if status == 200 else ""


# ---------------------------------------------------------------------------
# The journey
# ---------------------------------------------------------------------------

def test_real_document_ingestion(server, journey_app):
    report = {"journey": "GJ-06-REAL-DOCUMENTS", "steps": []}

    def step(name, ok, detail=""):
        report["steps"].append({"step": name, "ok": bool(ok), "detail": detail})

    http = Http(server.base)
    step("signin", bool(_signin(http)), "signed in")

    # ── 1. REAL PDF whose hierarchy can only come from its CONTENT ──────────
    pdf = make_pdf(BALI_ITINERARY_TEXT)
    step("pdf_bytes_are_real", pdf[:5] == b"%PDF-", f"magic={pdf[:5]!r} size={len(pdf)}")

    st, body = http.upload("/api/v1/documents/upload", NEUTRAL_PDF_NAME,
                           pdf, "application/pdf")
    step("pdf_uploaded", st == 201, f"POST upload -> {st}")
    doc = body.get("document") or body.get("data") or {}
    doc_id = doc.get("id") or doc.get("document_id")
    step("pdf_has_id", bool(doc_id), f"doc_id={doc_id}")

    intel = doc.get("intelligence") or doc.get("intelligence_summary") or {}
    text = doc.get("extracted_text") or intel.get("extracted_text") or ""

    step("pdf_text_extracted", "Bali" in text and len(text) > 80,
         f"extracted {len(text)} chars; contains Bali={'Bali' in text}")

    classification = doc.get("classification") or intel.get("classification")
    step("pdf_classified_from_content", classification == "itinerary",
         f"classification={classification!r}")

    hierarchy = doc.get("hierarchy") or intel.get("hierarchy") or []
    step("pdf_hierarchy_content_derived_len", len(hierarchy) >= 2,
         f"hierarchy={hierarchy}")
    if hierarchy:
        step("pdf_hierarchy_root", hierarchy[0] == "Itinerary",
             f"hierarchy[0]={hierarchy[0]!r}")
    if len(hierarchy) >= 2:
        step("pdf_hierarchy_scope_international", hierarchy[1] == "International",
             f"hierarchy[1]={hierarchy[1]!r}")
    step("pdf_hierarchy_names_destination",
         any("Bali" in str(seg) for seg in hierarchy),
         f"hierarchy={hierarchy}")

    entities = doc.get("entities") or intel.get("entities") or {}
    step("pdf_entities_extracted", bool(entities),
         f"entity keys={list(entities)[:6] if isinstance(entities, dict) else type(entities).__name__}")

    # ── 2. REAL DOCX quotation ─────────────────────────────────────────────
    docx_bytes = make_docx(QUOTATION_TEXT)
    step("docx_bytes_are_real", docx_bytes[:2] == b"PK", f"magic={docx_bytes[:2]!r}")

    st, body = http.upload("/api/v1/documents/upload", NEUTRAL_DOCX_NAME,
                           docx_bytes,
                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    step("docx_uploaded", st == 201, f"POST upload -> {st}")
    d2 = body.get("document") or body.get("data") or {}
    i2 = d2.get("intelligence") or d2.get("intelligence_summary") or {}
    t2 = d2.get("extracted_text") or i2.get("extracted_text") or ""
    c2 = d2.get("classification") or i2.get("classification")
    step("docx_text_extracted", "Quotation" in t2 or "quotation" in t2,
         f"extracted {len(t2)} chars")
    step("docx_classified_from_content", c2 == "quotation",
         f"classification={c2!r}")

    # ── 3. REAL XLSX ───────────────────────────────────────────────────────
    xlsx = make_xlsx()
    step("xlsx_bytes_are_real", xlsx[:2] == b"PK", f"magic={xlsx[:2]!r}")
    st, body = http.upload(
        "/api/v1/documents/upload", NEUTRAL_XLSX_NAME, xlsx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    step("xlsx_uploaded", st == 201, f"POST upload -> {st}")
    d3 = body.get("document") or body.get("data") or {}
    i3 = d3.get("intelligence") or d3.get("intelligence_summary") or {}
    t3 = d3.get("extracted_text") or i3.get("extracted_text") or ""
    step("xlsx_text_extracted", len(t3) > 0, f"extracted {len(t3)} chars")

    # ── 4. Persistence + restart ───────────────────────────────────────────
    st, listing = http.json("GET", "/api/v1/documents/")
    names = [d.get("filename") for d in (listing.get("documents") or listing.get("data") or [])]
    step("pdf_listed_after_upload", NEUTRAL_PDF_NAME in names,
         f"listed {len(names)} documents")

    st, fetched = http.json("GET", f"/api/v1/documents/{doc_id}")
    fdoc = fetched.get("document") or fetched.get("data") or {}
    step("pdf_re_fetch_keeps_classification",
         (fdoc.get("classification") == "itinerary"), 
         f"classification on re-fetch={fdoc.get('classification')!r}")

    restarted = Server(journey_app).start()
    try:
        h2 = Http(restarted.base)
        step("restart_signin", bool(_signin(h2)), "signed in after restart")
        st, l2 = h2.json("GET", "/api/v1/documents/")
        names2 = [d.get("filename") for d in (l2.get("documents") or l2.get("data") or [])]
        step("pdf_survives_restart", NEUTRAL_PDF_NAME in names2,
             f"after restart -> {st}")
        st, f2 = h2.json("GET", f"/api/v1/documents/{doc_id}")
        fd2 = f2.get("document") or f2.get("data") or {}
        step("classification_survives_restart",
             fd2.get("classification") == "itinerary",
             f"classification={fd2.get('classification')!r}")
    finally:
        restarted.stop()

    # ── 5. Tenant isolation ────────────────────────────────────────────────
    other = Http(server.base)
    step("other_tenant_signin", bool(_signin(other, EMAIL_ALT)), "alt org signed in")
    st, _ = other.json("GET", f"/api/v1/documents/{doc_id}")
    step("cross_tenant_document_denied", st == 404, f"other org GET -> {st}")
    st, l3 = other.json("GET", "/api/v1/documents/")
    n3 = [d.get("filename") for d in (l3.get("documents") or l3.get("data") or [])]
    step("cross_tenant_not_listed", NEUTRAL_PDF_NAME not in n3,
         f"other org list -> {st}")

    # ── 6. IMAGE — recorded honestly, not claimed ──────────────────────────
    png = make_png()
    st, body = http.upload("/api/v1/documents/upload", NEUTRAL_PNG_NAME,
                           png, "image/png")
    step("image_upload_accepted", st == 201, f"POST upload -> {st}")
    d4 = body.get("document") or body.get("data") or {}
    i4 = d4.get("intelligence") or d4.get("intelligence_summary") or {}
    t4 = d4.get("extracted_text") or i4.get("extracted_text") or ""
    report["image_ocr_note"] = {
        "extracted_chars": len(t4),
        "classification": d4.get("classification") or i4.get("classification"),
        "note": ("no OCR provider installed (pytesseract absent), so an image "
                 "cannot be content-classified in this environment"),
    }

    # ── Report ─────────────────────────────────────────────────────────────
    report["failed_steps"] = [s["step"] for s in report["steps"] if not s["ok"]]
    report["ok"] = not report["failed_steps"]
    out_dir = os.environ.get("JOURNEY_REPORT_DIR", tempfile.gettempdir())
    try:
        with open(os.path.join(out_dir, "journey_gj06_real_docs_report.json"),
                  "w", encoding="utf-8") as fh:
            json.dump(report, fh, indent=2)
    except OSError:
        pass

    assert report["ok"], (
        "GJ-06 real-document ingestion failed at: "
        + ", ".join(report["failed_steps"]) + "\n"
        + json.dumps(report["steps"], indent=2)
    )